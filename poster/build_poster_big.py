# -*- coding: utf-8 -*-
"""MailAgent 海报 · 大字版（所有文字 ≥ 44pt，内容精简）。

独立于 build_poster.py，输出单独文件，不影响原海报。
用法：
    uv run --with python-pptx python3 poster/build_poster_big.py
导出：
    cd poster
    soffice --headless --convert-to pdf 课程poster-MailAgent-大字版.pptx --outdir .
    pdftoppm -jpeg -r 200 -jpegopt quality=90 课程poster-MailAgent-大字版.pdf out && mv out-1.jpg 课程poster-MailAgent-大字版.jpg
"""

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "课程poster-模板.pptx")
OUT = os.path.join(HERE, "课程poster-MailAgent-大字版.pptx")

PRIMARY = RGBColor(0x8E, 0x5A, 0xA0)
PRIMARY2 = RGBColor(0xAD, 0x86, 0xC2)
NEUTRAL = RGBColor(0x94, 0x8C, 0xA0)
LIGHT = RGBColor(0xF1, 0xEA, 0xF6)
MID = RGBColor(0xE3, 0xD6, 0xEE)
BORDER = RGBColor(0xDD, 0xD0, 0xE8)
GREY = RGBColor(0x52, 0x4B, 0x5A)
CARDBG = RGBColor(0xF9, 0xF6, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CN = "微软雅黑"

MINPT = 44  # 最小字号


def set_font(run, size, color, bold=False, name=CN):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", name)


def para(tf, text, size, color, *, bold=False, sa=2, align=PP_ALIGN.CENTER, first=False):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align; p.space_after = Pt(sa); p.space_before = Pt(0)
    r = p.add_run(); r.text = text
    set_font(r, size, color, bold=bold)
    return p


def box(slide, x, y, w, h, lines, fill, tcolor, size=MINPT, *, bold=True,
        rounded=True, line_color=None):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shp.adjustments[0] = 0.10
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line_color is not None:
        shp.line.color.rgb = line_color; shp.line.width = Pt(1.5)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Inches(0.05))
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        para(tf, ln, size, tcolor, bold=bold, first=(i == 0))
    return shp


def arrow(slide, kind, x, y, w, h, color=NEUTRAL):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    shp.line.fill.background(); shp.shadow.inherit = False
    return shp


AR = MSO_SHAPE.RIGHT_ARROW
AD = MSO_SHAPE.DOWN_ARROW


def card(slide, x, y, w, h, title):
    body = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(x), Inches(y), Inches(w), Inches(h))
    body.adjustments[0] = 0.02
    body.fill.solid(); body.fill.fore_color.rgb = CARDBG
    body.line.color.rgb = BORDER; body.line.width = Pt(1.5)
    body.shadow.inherit = False
    hh = 1.45
    head = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(x), Inches(y), Inches(w), Inches(hh))
    head.adjustments[0] = 0.10
    head.fill.solid(); head.fill.fore_color.rgb = PRIMARY
    head.line.fill.background(); head.shadow.inherit = False
    htf = head.text_frame; htf.vertical_anchor = MSO_ANCHOR.MIDDLE
    htf.margin_left = Inches(0.35)
    para(htf, title, 50, WHITE, bold=True, align=PP_ALIGN.LEFT, first=True)
    return (x + 0.45, y + hh + 0.3, w - 0.9, h - hh - 0.6)


def vstack(body, items, box_h=1.7, fs=MINPT, gap=0.75):
    """卡片内竖向流程：整宽大框 + 向下箭头，垂直居中。"""
    bx, by, bw, bh = body
    n = len(items)
    total = n * box_h + (n - 1) * gap
    y = by + max(0, (bh - total) / 2)
    for i, it in enumerate(items):
        lines, fill, tc = it
        box(slide, bx, y, bw, box_h, lines, fill, tc, fs)
        if i < n - 1:
            arrow(slide, AD, bx + bw / 2 - 0.45, y + box_h + 0.04, 0.9, gap - 0.12)
        y += box_h + gap


# ───────── 组装 ─────────
prs = Presentation(SRC)
slide = prs.slides[0]

for sh in list(slide.shapes):
    if sh.has_text_frame and "建议格式" in sh.text_frame.text:
        sh._element.getparent().remove(sh._element)

for sh in slide.shapes:
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text
    if t.startswith("联系人"):
        sh.text_frame.paragraphs[0].runs[0].text = "作者：张翔宇  2022012081"
        sh.left = Inches(1.0); sh.width = Inches(33.43)
        for p in sh.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
    elif t.startswith("邮箱"):
        sh.text_frame.paragraphs[0].runs[0].text = "邮箱：xiangyuz22@mails.tsinghua.edu.cn"
        sh.left = Inches(1.0); sh.width = Inches(33.43)
        for p in sh.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER

# 项目标题
ptb = slide.shapes.add_textbox(Inches(0.6), Inches(5.0), Inches(34.2), Inches(2.6))
para(ptb.text_frame, "智能邮件管理代理  ·  Mail Agent", 66, PRIMARY, bold=True, sa=18, first=True)
para(ptb.text_frame, "大模型驱动的 Gmail + Google Calendar 邮件自动化代理", 46, GREY)

def hflow(body, items, desc, box_h=2.3, fs=MINPT, desc_fs=MINPT):
    """卡片内横向流程：整排大框 + 右向箭头，下方左对齐说明；整体垂直居中。"""
    bx, by, bw, bh = body
    n = len(items); aw = 0.9
    cw = (bw - (n - 1) * aw) / n
    group_h = box_h + 0.8 + 1.6
    top = by + max(0.0, (bh - group_h) / 2)
    for i, (lines, fill, tc) in enumerate(items):
        px = bx + i * (cw + aw)
        box(slide, px, top, cw, box_h, lines, fill, tc, fs)
        if i < n - 1:
            arrow(slide, AR, px + cw + 0.05, top + box_h / 2 - 0.3, aw - 0.1, 0.6)
    dy = top + box_h + 0.8
    tb = slide.shapes.add_textbox(Inches(bx), Inches(dy),
                                  Inches(bw), Inches(by + bh - dy - 0.05))
    tb.text_frame.word_wrap = True
    tb.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    para(tb.text_frame, desc, desc_fs, GREY, align=PP_ALIGN.LEFT, first=True)


# 架构（整宽：四步横排 + 分发 + 说明；加高、纵向均匀分布、间隙加大）
AX, AY, AW, AH = 0.7, 8.4, 34.03, 8.4
bx, by, bw, bh = card(slide, AX, AY, AW, AH, "系统架构")
step = [["Gmail API", "/ IMAP"], ["邮件解析", "Parser"],
        ["LLM", "语义分析"], ["意图路由", "Router"]]
n = len(step); g = 0.7
sw = (bw - (n - 1) * g) / n
row_y = by + 0.3
for i, t in enumerate(step):
    px = bx + i * (sw + g)
    box(slide, px, row_y, sw, 1.7, t, PRIMARY if i % 2 == 0 else PRIMARY2, WHITE, 44)
    if i < n - 1:
        arrow(slide, AR, px + sw + 0.05, row_y + 0.55, g - 0.1, 0.6)
# 加大的向下分发箭头
arrow(slide, AD, bx + bw / 2 - 0.5, row_y + 1.9, 1.0, 0.95)
disp_y = row_y + 2.9
box(slide, bx, disp_y, bw, 1.5, "分发 → 日历调度 · 回复生成 · 优先级排序",
    LIGHT, PRIMARY, 44)
adesc_y = disp_y + 2.3
adesc = slide.shapes.add_textbox(Inches(bx), Inches(adesc_y),
                                 Inches(bw), Inches(by + bh - adesc_y - 0.1))
adesc.text_frame.word_wrap = True
adesc.text_frame.vertical_anchor = MSO_ANCHOR.TOP
para(adesc.text_frame,
     "拉取 → 大模型理解 → 意图分发到 日历 / 回复 / 优先级；来源与日历后端均可替换。",
     44, GREY, align=PP_ALIGN.LEFT, first=True)

# 六张功能卡片（横向流程 + 下方说明）
LX, RX, CW = 0.7, 18.03, 16.7
ROWY = [17.0, 25.6, 34.2]
CH = 8.3

b = card(slide, LX, ROWY[0], CW, CH, "邮件语义理解")
hflow(b, [(["原始", "邮件"], PRIMARY, WHITE),
          (["GLM-4.6", "语义分析"], PRIMARY2, WHITE),
          (["意图·紧急度", "需回复·含日程"], LIGHT, PRIMARY)],
      "用大模型读懂邮件：判别意图、紧急度、是否需回复与是否含日程，供后续模块复用。")

b = card(slide, RX, ROWY[0], CW, CH, "优先级评分")
hflow(b, [(["发件人·截止", "意图·紧急"], LIGHT, PRIMARY),
          (["综合", "评分"], PRIMARY, WHITE),
          (["Critical/High", "Medium/Low"], PRIMARY2, WHITE)],
      "结合发件人权重、截止日期与 LLM 分析综合打分，重要邮件自动排到最前。")

b = card(slide, LX, ROWY[1], CW, CH, "日历调度 · 冲突避让")
hflow(b, [(["解析", "起止时间"], PRIMARY2, WHITE),
          (["冲突", "检测"], PRIMARY2, WHITE),
          (["建事件 /", "避让改期"], PRIMARY, WHITE)],
      "解析自然语言时间并检测冲突；冲突时按日历空闲就近避让或建议改期。")

b = card(slide, RX, ROWY[1], CW, CH, "个性化回复生成")
hflow(b, [(["原邮件 +", "用户风格"], LIGHT, PRIMARY),
          (["GLM-4.6", "生成回复"], PRIMARY2, WHITE),
          (["Gmail", "草稿/发送"], PRIMARY, WHITE)],
      "结合邮件内容与你的写作风格生成自然回复；冲突时引用空闲建议改期，可存草稿或发送。")

b = card(slide, LX, ROWY[2], CW, CH, "Google 生态集成")
hflow(b, [(["Mail Agent", "OAuth2"], PRIMARY, WHITE),
          (["Gmail API", "读/发/标签"], PRIMARY2, WHITE),
          (["Calendar", "API · 日程"], PRIMARY2, WHITE)],
      "OAuth2 一次授权，原生读写真实 Gmail 与 Google Calendar，无需导出或中转。")

b = card(slide, RX, ROWY[2], CW, CH, "技术栈")
hflow(b, [(["GLM-4.6", "大模型"], PRIMARY, WHITE),
          (["FastAPI", "后端"], PRIMARY2, WHITE),
          (["HTMX 前端", "Typer CLI"], MID, PRIMARY)],
      "大模型 GLM-4.6 + FastAPI 后端 + HTMX 前端 / Typer 命令行，pytest 全覆盖。")

# 收尾：补东亚字体 + 把任何显式 <44pt 的字号提升到 44
for sh in slide.shapes:
    if not sh.has_text_frame:
        continue
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            rPr = r._r.get_or_add_rPr()
            for tag in ("a:ea", "a:cs"):
                el = rPr.find(qn(tag))
                if el is None:
                    el = rPr.makeelement(qn(tag), {}); rPr.append(el)
                el.set("typeface", CN)
            if r.font.size is not None and r.font.size < Pt(MINPT):
                r.font.size = Pt(MINPT)

prs.save(OUT)
print("已生成:", OUT, "| 形状数:", len(slide.shapes))
