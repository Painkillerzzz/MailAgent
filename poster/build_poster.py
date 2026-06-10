# -*- coding: utf-8 -*-
"""MailAgent 功能海报生成脚本：顶部架构图 + 6 张「每框带图」的功能卡片。
仅介绍功能，不含任何测试/指标内容。

用法（任意目录均可）：
    uv run --with python-pptx python3 poster/build_poster.py

导出 PDF / 高清 JPG（需 libreoffice + poppler-utils）：
    cd poster
    soffice --headless --convert-to pdf 课程poster-MailAgent.pptx --outdir .
    pdftoppm -jpeg -r 200 -jpegopt quality=90 课程poster-MailAgent.pdf out && mv out-1.jpg 课程poster-MailAgent.jpg

手动改要点：
    - 颜色：改下方 PRIMARY / PRIMARY2 / LIGHT / MID 等常量
    - 文案/联系人：见各 card(...) 调用与「联系人/邮箱」替换处
    - 版面：ROWY（三行 Y 坐标）、CH（卡片高度）、LX/RX/CW（左右列与宽度）
    - 各卡片小图在对应「# 1. ...」注释段里调整
"""

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# 路径基于脚本所在目录（poster/），可在任意 cwd 运行
HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "课程poster-模板.pptx")
OUT = os.path.join(HERE, "课程poster-MailAgent.pptx")

# 克制紫色（与品牌带 #8C3A92 同族）
PRIMARY = RGBColor(0x8E, 0x5A, 0xA0)
PRIMARY2 = RGBColor(0xAD, 0x86, 0xC2)
NEUTRAL = RGBColor(0x94, 0x8C, 0xA0)
LIGHT = RGBColor(0xF1, 0xEA, 0xF6)
MID = RGBColor(0xE3, 0xD6, 0xEE)
BORDER = RGBColor(0xDD, 0xD0, 0xE8)
DARK = RGBColor(0x2A, 0x24, 0x30)
GREY = RGBColor(0x52, 0x4B, 0x5A)
CARDBG = RGBColor(0xF9, 0xF6, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CN = "微软雅黑"

# 正文(卡片/图/说明)统一放大系数；标题与联系人不走这里，故不受影响
BODY_FS = 1.16


def set_font(run, size, color, bold=False, name=CN):
    run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = color; run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", name)


def para(tf, text, size, color, *, bold=False, sa=2, align=PP_ALIGN.CENTER, first=False):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align; p.space_after = Pt(sa); p.space_before = Pt(0)
    r = p.add_run(); r.text = text
    set_font(r, size, color, bold=bold)
    return p


def box(slide, x, y, w, h, lines, fill, tcolor, size, *, bold=True,
        rounded=True, line_color=None, lw=1.0):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shp.adjustments[0] = 0.14
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line_color is not None:
        shp.line.color.rgb = line_color; shp.line.width = Pt(lw)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Inches(0.04))
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        para(tf, ln, size * BODY_FS, tcolor, bold=bold, first=(i == 0))
    return shp


def arrow(slide, kind, x, y, w, h, color=NEUTRAL):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    shp.line.fill.background(); shp.shadow.inherit = False
    return shp


AR = MSO_SHAPE.RIGHT_ARROW
AD = MSO_SHAPE.DOWN_ARROW


def card(slide, x, y, w, h, title):
    body = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(x), Inches(y), Inches(w), Inches(h))
    body.adjustments[0] = 0.02
    body.fill.solid(); body.fill.fore_color.rgb = CARDBG
    body.line.color.rgb = BORDER; body.line.width = Pt(1.5)
    body.shadow.inherit = False
    hh = 1.0
    head = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(x), Inches(y), Inches(w), Inches(hh))
    head.adjustments[0] = 0.12
    head.fill.solid(); head.fill.fore_color.rgb = PRIMARY
    head.line.fill.background(); head.shadow.inherit = False
    htf = head.text_frame; htf.vertical_anchor = MSO_ANCHOR.MIDDLE
    htf.margin_left = Inches(0.3)
    para(htf, title, 28 * BODY_FS, WHITE, bold=True, align=PP_ALIGN.LEFT, first=True)
    return (x + 0.5, y + hh + 0.35, w - 1.0, h - hh - 0.7)  # body rect


def hcaption(slide, x, y, w, text, h=1.4):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    para(tf, text, 23 * BODY_FS, GREY, align=PP_ALIGN.CENTER, first=True)


def note(slide, x, y, w, h, text, size=21):
    """左对齐多行解释性文字（用于卡片内的说明段落）。"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    para(tf, text, size * BODY_FS, GREY, align=PP_ALIGN.LEFT, sa=3, first=True)


# ───────── 组装 ─────────
prs = Presentation(SRC)
slide = prs.slides[0]

for sh in list(slide.shapes):
    if sh.has_text_frame and "建议格式" in sh.text_frame.text:
        sh._element.getparent().remove(sh._element)

for sh in slide.shapes:
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text
    if t.startswith("联系人") or t.startswith("邮箱"):
        sh.text_frame.paragraphs[0].runs[0].text = (
            "作者：张翔宇  2022012081" if t.startswith("联系人")
            else "邮箱：xiangyuz22@mails.tsinghua.edu.cn")
        # 加宽到整幅并居中，避免长邮箱换行
        sh.left = Inches(1.0); sh.width = Inches(33.43)
        sh.text_frame.word_wrap = True
        for p in sh.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER

ptb = slide.shapes.add_textbox(Inches(0.6), Inches(5.45), Inches(34.2), Inches(2.2))
para(ptb.text_frame, "智能邮件管理代理  ·  Mail Agent", 60, PRIMARY, bold=True, sa=20, first=True)
para(ptb.text_frame, "大模型驱动的 Gmail + Google Calendar 邮件自动化代理", 30, GREY)

# ===== 图1：系统架构（整宽）=====
AX, AY, AW, AH = 0.7, 8.1, 34.03, 7.95
bx, by, bw, bh = card(slide, AX, AY, AW, AH, "系统架构  System Architecture")
# 第一行 5 个框：…→意图路由→意图分发（意图分发与其箭头同在此行）
row1 = [
    ("Gmail API / IMAP", PRIMARY, WHITE),
    ("邮件解析\nParser", PRIMARY2, WHITE),
    ("LLM 语义分析\nGLM-4.6", PRIMARY, WHITE),
    ("意图路由\nRouter", PRIMARY2, WHITE),
    ("意图分发\nDispatch", LIGHT, PRIMARY),
]
gap = 0.8
rbw = (bw - 4 * gap) / 5
rbh = 1.7
y1 = by + 0.2
for i, (txt, f, tc) in enumerate(row1):
    px = bx + i * (rbw + gap)
    box(slide, px, y1, rbw, rbh, txt.split("\n"), f, tc, 19)
    if i < 4:
        arrow(slide, AR, px + rbw + 0.08, y1 + rbh / 2 - 0.32, gap - 0.16, 0.64)
# 由「意图分发」向下分发到三类输出
y2 = y1 + rbh + 1.45
outs = ["Google Calendar\n日程事件", "Gmail 草稿 / 发送", "优先级评分\n收件箱排序"]
ow, ogap = 9.6, 1.1
ox = bx + (bw - (3 * ow + 2 * ogap)) / 2
for i, txt in enumerate(outs):
    px = ox + i * (ow + ogap)
    arrow(slide, AD, px + ow / 2 - 0.35, y2 - 0.85, 0.7, 0.78)
    box(slide, px, y2, ow, 1.6, txt.split("\n"), PRIMARY2, WHITE, 21)
note(slide, bx + 0.2, y2 + 2.15, bw - 0.4, AY + AH - (y2 + 2.15) - 0.2,
     "端到端流水线：从邮箱（Gmail API 或 IMAP）拉取未读邮件，经大模型完成语义理解后，由意图路由"
     "识别邮件类型并分发到三条下游——日历调度（解析时间、检测冲突、写入 Google Calendar）、回复"
     "生成（按你的风格生成草稿或发送）、优先级排序（综合发件人与截止日期打分）。各模块分层解耦、"
     "可独立测试；邮件来源（IMAP / Gmail API）与日历后端（本地 JSON / Google Calendar）均可替换。",
     size=18.5)

# ===== 6 张功能卡片（每框带图）=====
LX, RX, CW = 0.7, 18.03, 16.7
ROWY = [16.2, 25.4, 34.6]
CH = 8.9


def step_chain(items, bx, by, bw, bh, sizes=None):
    """在 body 内水平排列若干框 + 箭头，自动均分。items: [(lines, fill, tcolor)]"""
    n = len(items)
    aw = 1.0
    total_arrow = aw * (n - 1)
    cw = (bw - total_arrow) / n
    h = min(bh, 2.4)
    yy = by + (bh - h) / 2
    for i, (lines, fill, tc) in enumerate(items):
        px = bx + i * (cw + aw)
        box(slide, px, yy, cw, h, lines, fill, tc, 18)
        if i < n - 1:
            arrow(slide, AR, px + cw + 0.12, yy + h / 2 - 0.32, aw - 0.24, 0.64)


# 1. 邮件语义理解：邮件 → LLM → 4 个标签（图在上，说明在下）
b = card(slide, LX, ROWY[0], CW, CH, "邮件语义理解")
bx, by, bw, bh = b
dy = by + 0.3
box(slide, bx, dy, 3.2, 2.2, ["原始邮件"], PRIMARY, WHITE, 19)
arrow(slide, AR, bx + 3.35, dy + 0.78, 0.9, 0.64)
box(slide, bx + 4.4, dy, 3.6, 2.2, ["GLM-4.6", "语义分析"], PRIMARY2, WHITE, 18)
arrow(slide, AR, bx + 8.15, dy + 0.78, 0.9, 0.64)
tags = ["意图", "紧急度", "需回复", "含日程"]
tx = bx + 9.25
for i, t in enumerate(tags):
    cxx = tx + (i % 2) * 3.3
    cyy = dy - 0.4 + (i // 2) * 1.7
    box(slide, cxx, cyy, 3.0, 1.4, t, LIGHT, PRIMARY, 18, line_color=PRIMARY2)
note(slide, bx, dy + 3.1, bw, 1.9,
     "系统用大模型读懂每封邮件，而非简单关键词匹配：判断邮件意图（会议 / 任务 / 通知等）、"
     "紧急程度、是否需要你亲自回复、以及是否包含日程信息。它能结合语气与上下文做判断——例如"
     "区分『仅供知悉』的通知与真正需要你行动的请求，输出结构化结果供后续模块复用。")
# 分析示例（左对齐，展示真实抽取结果）
axb = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(bx), Inches(by + 5.05),
    Inches(bw), Inches(bh - 5.05 - 0.2))
axb.adjustments[0] = 0.05
axb.fill.solid(); axb.fill.fore_color.rgb = LIGHT
axb.line.color.rgb = PRIMARY2; axb.line.width = Pt(1.25)
axb.shadow.inherit = False
atf = axb.text_frame
atf.word_wrap = True
atf.vertical_anchor = MSO_ANCHOR.TOP
atf.margin_left = Inches(0.3); atf.margin_right = Inches(0.25)
atf.margin_top = Inches(0.14)
para(atf, "分析示例", 18 * BODY_FS, PRIMARY, bold=True, align=PP_ALIGN.LEFT, sa=12, first=True)
para(atf, "邮件：「李教授：周四下午 3 点方便讨论论文实验吗？」", 17 * BODY_FS, GREY,
     align=PP_ALIGN.LEFT, sa=12)
para(atf, "→ 意图＝会议请求 · 紧急度＝中 · 需回复＝是 · 含日程＝是（周四 15:00）",
     17 * BODY_FS, PRIMARY, align=PP_ALIGN.LEFT, sa=0)

# 2. 优先级评分：4 因子 → 评分 → 等级条
b = card(slide, RX, ROWY[0], CW, CH, "优先级评分")
bx, by, bw, bh = b
facs = ["发件人权重", "截止日期", "邮件意图", "紧急程度"]
fy = by + 0.1
ph = 2.5  # 评分/等级区垂直中心参考
for i, f in enumerate(facs):
    box(slide, bx, fy + i * 1.35, 4.6, 1.2, f, LIGHT, PRIMARY, 17, line_color=PRIMARY2)
arrow(slide, AR, bx + 4.75, by + ph - 0.4, 1.1, 0.8)
box(slide, bx + 6.0, by + ph - 1.0, 4.0, 2.0, ["综合", "评分"], PRIMARY, WHITE, 20)
arrow(slide, AR, bx + 10.15, by + ph - 0.4, 1.1, 0.8)
levels = [("Critical", PRIMARY), ("High", RGBColor(0x9E, 0x6E, 0xB4)),
          ("Medium", PRIMARY2), ("Low", MID)]
ly = by + ph - 1.7
for i, (lv, c) in enumerate(levels):
    tc = WHITE if i < 3 else PRIMARY
    box(slide, bx + 11.25, ly + i * 0.95, 4.2, 0.82, lv, c, tc, 16)
hcaption(slide, bx, by + 5.9, bw,
         "规则(发件人权重 / 截止日期)与 LLM 综合打分，重要邮件自动排到最前")

# 3. 日历调度 · 冲突避让：解析 → 冲突检测 →（空闲感知）建议/避让改期
b = card(slide, LX, ROWY[1], CW, CH, "日历调度 · 冲突避让")
bx, by, bw, bh = b
# 第一行：时间短语 → LLM 解析 → 冲突检测
w1 = (bw - 2 * 0.7) / 3
h1 = 1.7
chain = [
    (["时间短语", "「周四 15:00」"], LIGHT, PRIMARY),
    (["LLM 解析", "起止时间"], PRIMARY2, WHITE),
    (["冲突检测"], PRIMARY2, WHITE),
]
for i, (t, f, tc) in enumerate(chain):
    px = bx + i * (w1 + 0.7)
    box(slide, px, by + 0.2, w1, h1, t, f, tc, 17)
    if i < 2:
        arrow(slide, AR, px + w1 + 0.14, by + 0.2 + h1 / 2 - 0.3, 0.7 - 0.28, 0.6)
# 第二行：两种结果（高亮空闲感知改期）
ry = by + 0.2 + h1 + 1.05
w2 = (bw - 0.6) / 2
outs = [
    (["无冲突", "→ 建 Google Calendar 事件"], PRIMARY, WHITE),
    (["有冲突 → 查日历空闲", "避让 / 建议最近可行时段"], PRIMARY2, WHITE),
]
for i, (t, f, tc) in enumerate(outs):
    px = bx + i * (w2 + 0.6)
    arrow(slide, AD, px + w2 / 2 - 0.3, ry - 0.85, 0.6, 0.65)
    box(slide, px, ry, w2, 1.8, t, f, tc, 17)
note(slide, bx, ry + 2.7, bw, bh - (ry - by) - 2.9,
     "从邮件中解析出具体起止时间，并与日历中已有事件比对。无冲突则直接创建 Google Calendar "
     "事件；一旦冲突，会在工作时间内查找最近的空闲时段，自动避让，或在回复中给出可行的改期建议。")

# 4. 个性化回复：原邮件 + 用户风格 → LLM → 草稿 → 发送
b = card(slide, RX, ROWY[1], CW, CH, "个性化回复生成")
bx, by, bw, bh = b
box(slide, bx, by + 0.2, 4.4, 1.5, "原邮件", LIGHT, PRIMARY, 18, line_color=PRIMARY2)
box(slide, bx, by + 2.0, 4.4, 1.5, "用户风格 / 签名", LIGHT, PRIMARY, 17, line_color=PRIMARY2)
arrow(slide, AR, bx + 4.55, by + 1.35, 1.0, 0.8)
box(slide, bx + 5.7, by + 0.7, 4.2, 2.2, ["GLM-4.6", "生成回复"], PRIMARY2, WHITE, 18)
arrow(slide, AR, bx + 10.05, by + 1.35, 1.0, 0.8)
box(slide, bx + 11.2, by + 0.7, 4.5, 2.2, ["回复草稿", "Gmail 草稿/发送"], PRIMARY, WHITE, 18)
note(slide, bx, by + 3.85, bw, 1.35,
     "结合原邮件内容与你的写作风格、签名，用大模型生成自然、得体的回复草稿；若与现有日程冲突，"
     "会引用日历空闲时段提出具体的改期时间，确认后一键存草稿或发送。")
# 回复草稿示例（左对齐，多个场景示例）
exb = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(bx), Inches(by + 5.0),
    Inches(bw), Inches(bh - 5.0 - 0.2))
exb.adjustments[0] = 0.05
exb.fill.solid(); exb.fill.fore_color.rgb = LIGHT
exb.line.color.rgb = PRIMARY2; exb.line.width = Pt(1.25)
exb.shadow.inherit = False
etf = exb.text_frame
etf.word_wrap = True
etf.vertical_anchor = MSO_ANCHOR.TOP
etf.margin_left = Inches(0.3); etf.margin_right = Inches(0.25)
etf.margin_top = Inches(0.14)
para(etf, "回复草稿示例", 18 * BODY_FS, PRIMARY, bold=True, align=PP_ALIGN.LEFT, sa=4, first=True)
for line in [
    "· 会议冲突：「周四 15:00 与组会冲突，建议改到 16:00 或周五 10:00，您看是否方便？」",
    "· 任务确认：「收到，我会在本周五前完成季度报告并发您审阅。」",
    "· 礼貌婉拒：「感谢邀请，这周时间已排满，下周二之后我都方便。」",
]:
    para(etf, line, 17 * BODY_FS, GREY, align=PP_ALIGN.LEFT, sa=3)

# 5. Google 生态集成：Agent ↔ Gmail API / Calendar API (OAuth2)
b = card(slide, LX, ROWY[2], CW, CH, "Google 生态集成")
bx, by, bw, bh = b
cx = bx + bw / 2 - 2.8
box(slide, bx + 0.2, by + 0.3, 6.0, 2.3,
    ["Gmail API", "读取 / 发送 / 标记已读", "草稿与标签管理"], PRIMARY2, WHITE, 17)
box(slide, bx + bw - 6.2, by + 0.3, 6.0, 2.3,
    ["Calendar API", "日程读写 / 查询", "冲突检测"], PRIMARY2, WHITE, 17)
box(slide, cx, by + 2.95, 5.6, 1.9, ["Mail Agent", "OAuth2 一次授权"], PRIMARY, WHITE, 19)
arrow(slide, MSO_SHAPE.LEFT_ARROW, bx + 6.35, by + 1.35, 1.3, 0.7)
arrow(slide, AR, bx + bw - 7.65, by + 1.35, 1.3, 0.7)
note(slide, bx, by + 5.2, bw, bh - 5.5,
     "通过 OAuth2 一次授权，原生读写真实的 Gmail 与 Google Calendar（scopes：gmail.modify / "
     "gmail.send / calendar）：拉取未读、建 / 查日程并检测冲突、生成草稿或发送；未授权或离线时"
     "自动降级为本地模式，其余功能照常运行。")

# 6. 技术栈：分层
b = card(slide, RX, ROWY[2], CW, CH, "技术栈")
bx, by, bw, bh = b
layers = [
    ("大模型  GLM-4.6（Coding Plan 端点）", PRIMARY),
    ("后端  FastAPI + Pydantic v2", PRIMARY2),
    ("前端  HTMX + Alpine.js 仪表板", PRIMARY2),
    ("命令行  Typer + Rich", MID),
]
area = bh - 1.7          # 预留底部说明文字空间（含上下留白）
lh = (area - 0.3 * 3) / 4
for i, (t, c) in enumerate(layers):
    tc = WHITE if c != MID else PRIMARY
    box(slide, bx, by + i * (lh + 0.3), bw, lh, t, c, tc, 18, rounded=True)
hcaption(slide, bx, by + area + 0.45, bw,
         "全栈一体：大模型 + Web 后端 / 前端 + CLI，pytest 覆盖")

# 收尾：补东亚字体
for sh in slide.shapes:
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                rPr = r._r.get_or_add_rPr()
                for tag in ("a:ea", "a:cs"):
                    el = rPr.find(qn(tag))
                    if el is None:
                        el = rPr.makeelement(qn(tag), {}); rPr.append(el)
                    el.set("typeface", CN)

prs.save(OUT)
print("已生成:", OUT, "| 形状数:", len(slide.shapes))
